"""Local paragraph extractor and embedding worker. Original files are never modified."""
import json
import os
import re
import subprocess
import tempfile
import time
import urllib.error
import urllib.request
from pathlib import Path

import fitz
import psycopg
from docx import Document
from pptx import Presentation

DB = os.environ["DATABASE_URL"]
DATA = Path(os.environ.get("DATA_ROOT", "/data"))
OLLAMA = os.environ.get("OLLAMA_URL", "http://ollama:11434").rstrip("/")
DIRECT = {".pdf", ".docx", ".pptx", ".txt", ".csv"}
CONVERT = {".doc", ".ppt", ".rtf", ".odt", ".odp", ".xls", ".xlsx", ".ods"}
SUPPORTED = DIRECT | CONVERT


def embed(texts, model_name, dimensions):
    request = urllib.request.Request(
        OLLAMA + "/api/embed",
        data=json.dumps({"model": model_name, "input": texts}).encode(),
        headers={"Content-Type": "application/json"},
    )
    with urllib.request.urlopen(request, timeout=300) as response:
        vectors = json.load(response)["embeddings"]
    if len(vectors) != len(texts) or any(len(vector) != dimensions for vector in vectors):
        raise RuntimeError("向量模型返回了错误的维度")
    return vectors


def clean(value):
    return re.sub(r"\s+", " ", value).strip()


def pdf_paragraphs(file_path):
    with fitz.open(file_path) as document:
        for page_number, page in enumerate(document, 1):
            for block in page.get_text("blocks"):
                if len(block) > 6 and block[6] != 0:
                    continue
                text = clean(block[4])
                if text:
                    yield f"第 {page_number} 页", text


def docx_paragraphs(file_path):
    document = Document(file_path)
    heading = ""
    for item in document.iter_inner_content():
        if hasattr(item, "paragraphs"):
            for row in item.rows:
                cells = [clean(cell.text) for cell in row.cells]
                text = " | ".join(cells)
                if text.strip(" |"):
                    yield heading or "表格", text
        else:
            text = clean(item.text)
            if not text:
                continue
            if item.style and item.style.name.lower().startswith("heading"):
                heading = text[:120]
            else:
                yield heading or "正文", text


def pptx_paragraphs(file_path):
    presentation = Presentation(file_path)
    for slide_number, slide in enumerate(presentation.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = clean(paragraph.text)
                    if text:
                        yield f"第 {slide_number} 页幻灯片", text
            if shape.has_table:
                for row in shape.table.rows:
                    text = " | ".join(clean(cell.text) for cell in row.cells)
                    if text.strip(" |"):
                        yield f"第 {slide_number} 页幻灯片", text


def text_paragraphs(file_path):
    # Incremental decoding bounds memory even for a very large plain text file.
    import codecs
    with open(file_path, "rb") as source:
        sample = source.read(8192)
        source.seek(0)
        encoding = "utf-8-sig"
        if sample.startswith(b"\xff\xfe"):
            encoding = "utf-16le"
        elif sample.startswith(b"\xfe\xff"):
            encoding = "utf-16be"
        else:
            try:
                sample.decode("utf-8-sig")
            except UnicodeDecodeError:
                encoding = "gb18030"
        decoder = codecs.getincrementaldecoder(encoding)(errors="replace")
        buffer = ""
        line_number = 1
        while block := source.read(65536):
            buffer += decoder.decode(block)
            while "\n" in buffer or len(buffer) > 2400:
                offset = buffer.find("\n")
                cut = offset + 1 if 0 <= offset < 2400 else 1200
                part, buffer = buffer[:cut], buffer[cut:]
                text = clean(part)
                if text:
                    yield f"第 {line_number} 行附近", text
                line_number += part.count("\n")
        buffer += decoder.decode(b"", final=True)
        if clean(buffer):
            yield f"第 {line_number} 行附近", clean(buffer)


def paragraphs(file_path, extension):
    if extension == ".pdf":
        yield from pdf_paragraphs(file_path)
    elif extension == ".docx":
        yield from docx_paragraphs(file_path)
    elif extension == ".pptx":
        yield from pptx_paragraphs(file_path)
    elif extension in {".txt", ".csv"}:
        yield from text_paragraphs(file_path)
    else:
        with tempfile.TemporaryDirectory(prefix="index-", dir=DATA / "previews") as folder:
            profile = Path(folder) / "profile"
            subprocess.run([
                "libreoffice", f"-env:UserInstallation=file://{profile}",
                "--headless", "--convert-to", "pdf", "--outdir", folder, str(file_path)
            ], check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=300)
            converted = Path(folder) / (file_path.stem + ".pdf")
            if not converted.is_file():
                raise RuntimeError("Office 转换未生成 PDF")
            yield from pdf_paragraphs(converted)


def chunks(parts):
    current = []
    current_location = ""
    length = 0
    for location, paragraph in parts:
        if not paragraph:
            continue
        if location != current_location and current:
            yield current_location, "\n".join(current)
            current, length = [], 0
        current_location = location
        # Keep an oversized paragraph readable and give adjacent windows context.
        segments = []
        if len(paragraph) <= 1200:
            segments = [paragraph]
        else:
            for offset in range(0, len(paragraph), 1050):
                segments.append(paragraph[offset:offset + 1200])
        for segment in segments:
            if current and length + len(segment) > 1200:
                yield current_location, "\n".join(current)
                current, length = [], 0
            current.append(segment)
            length += len(segment)
            if length >= 700:
                yield current_location, "\n".join(current)
                current, length = [], 0
    if current:
        yield current_location, "\n".join(current)


def index_one(conn, job):
    job_id, document_id, extension, model_name, dimensions = job
    file_path = DATA / "raw" / (document_id + extension)
    try:
        if not file_path.is_file():
            raise RuntimeError("原始文件不存在")
        if extension in CONVERT | {".docx", ".pptx"} and file_path.stat().st_size > 512 * 1024 ** 2:
            raise RuntimeError("Office 文件超过 512 MiB，已保存原件，暂不建立段落索引")
        with conn.cursor() as cursor:
            cursor.execute("DELETE FROM document_parse_chunks WHERE job_id=%s", (job_id,))
            cursor.execute("UPDATE document_parse_jobs SET status='processing', error=NULL, updated_at=now() WHERE id=%s", (job_id,))
        conn.commit()
        batch = []
        count = 0
        for locator, content in chunks(paragraphs(file_path, extension)):
            batch.append((locator, content))
            if len(batch) >= 16:
                count = save_batch(conn, job_id, count, batch, model_name, dimensions)
                batch.clear()
        if batch:
            count = save_batch(conn, job_id, count, batch, model_name, dimensions)
        if count == 0:
            raise RuntimeError("未提取到文字；扫描版 PDF 需要后续 OCR")
        with conn.cursor() as cursor:
            cursor.execute("UPDATE document_parse_jobs SET status='ready', chunk_count=%s, updated_at=now() WHERE id=%s", (count, job_id))
        conn.commit()
        print(f"Indexed {document_id}: {count} chunks", flush=True)
    except Exception as error:
        conn.rollback()
        temporary = isinstance(error, (urllib.error.URLError, TimeoutError))
        with conn.cursor() as cursor:
            cursor.execute("DELETE FROM document_parse_chunks WHERE job_id=%s", (job_id,))
            cursor.execute("UPDATE document_parse_jobs SET status=%s, error=%s, chunk_count=0, updated_at=now() WHERE id=%s",
                           ("pending" if temporary else "failed", str(error)[:500], job_id))
        conn.commit()
        print(f"Index failed {document_id}: {error}", flush=True)
        if temporary:
            time.sleep(10)


def save_batch(conn, job_id, offset, batch, model_name, dimensions):
    vectors = embed([content for _, content in batch], model_name, dimensions)
    with conn.cursor() as cursor:
        for number, ((locator, content), vector) in enumerate(zip(batch, vectors), offset):
            cursor.execute("INSERT INTO document_parse_chunks (job_id, chunk_index, locator, content, embedding) VALUES (%s,%s,%s,%s,%s::vector)",
                           (job_id, number, locator, content, "[" + ",".join(map(str, vector)) + "]"))
    conn.commit()
    return offset + len(batch)


def work(conn):
    with conn.cursor() as cursor:
        cursor.execute("UPDATE document_parse_jobs SET status='pending' WHERE status='processing'")
    conn.commit()
    while True:
        with conn.cursor() as cursor:
            cursor.execute("""SELECT j.id,j.source_document_id,j.extension,m.model_name,m.dimensions
                FROM document_parse_jobs j JOIN ai_models m ON m.id=j.embedding_model_id
                WHERE j.status='pending' ORDER BY j.created_at ASC LIMIT 1""")
            job = cursor.fetchone()
        conn.commit()
        if job:
            index_one(conn, job)
        else:
            time.sleep(5)


if __name__ == "__main__":
    while True:
        try:
            with psycopg.connect(DB, autocommit=False) as connection:
                work(connection)
        except Exception as exc:
            print(f"Indexer reconnecting: {exc}", flush=True)
            time.sleep(5)
