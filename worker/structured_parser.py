"""Local structure-preserving extraction and Unicode-character chunking.

All parsers emit content/content_type/source/section_path blocks. Source metadata
is retained through splitting; no LLM is involved in extraction or chunking.
"""
import csv
import io
import json
import re
import subprocess
import tempfile
from pathlib import Path

import fitz
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pptx import Presentation

MANIFEST = json.loads(Path(__file__).with_name("parser_manifest.json").read_text(encoding="utf-8"))
PARSER_VERSION = MANIFEST["parserVersion"]
DEFAULTS = MANIFEST["defaults"]
CONVERT = {".doc", ".ppt", ".rtf", ".odt", ".odp", ".xls", ".ods"}


def clean(value):
    return re.sub(r"[\t\r\f\v ]+", " ", str(value or "")).strip()


def cell_text(value):
    # Numeric zero and False must not be treated as empty cells.
    return "" if value is None else clean(str(value))


def block(content, kind, source, section=None):
    return {"content": content, "content_type": kind, "source": source, "section_path": list(section or [])}


def decode_text(data):
    if data.startswith(b"\xff\xfe") or data.startswith(b"\xfe\xff"):
        return data.decode("utf-16"), "utf-16"
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            pass
    raise ValueError("无法按 UTF-8、UTF-16 或 GB18030 解码，请转换为 UTF-8 后重试")


def table_blocks(rows, source, section, config, table_id):
    """rows is [(original_row_number, [cell values ...]), ...]."""
    if not rows:
        return []
    width = max(len(values) for _, values in rows)
    count = min(config["headerRows"], len(rows))
    headers = []
    for col in range(width):
        parts = list(dict.fromkeys(values[col] for _, values in rows[:count] if col < len(values) and values[col]))
        headers.append(" / ".join(parts) or f"列{col + 1}")
    # Keep duplicate column names distinguishable while retaining their units.
    if len(set(headers)) != len(headers):
        headers = [f"{name} [列{i + 1}]" for i, name in enumerate(headers)]
    result = []
    for number, values in rows[count:]:
        values = values + [""] * (width - len(values))
        if not any(values):
            continue
        content = " | ".join(f"{headers[i]}: {value if value else '∅'}" for i, value in enumerate(values))
        result.append(block(content, "table", {**source, "table_id": table_id, "row": number,
            "column_names": headers, "cells": values, "header_rows": [r[0] for r in rows[:count]]}, section))
    if not result:
        result.append(block(" | ".join(headers), "table", {**source, "table_id": table_id,
            "row": rows[0][0], "column_names": headers, "cells": [], "header_only": True}, section))
    return result


def extract_docx(path, config):
    doc, result, headings = Document(path), [], []
    pnum, tnum = 0, 0
    for item in doc.iter_inner_content():
        if isinstance(item, Paragraph):
            pnum += 1
            text = clean(item.text)
            if not text:
                continue
            style = item.style
            match = re.search(r"(?:heading|标题)\s*(\d+)", f"{style.name} {style.style_id}" if style else "", re.I)
            outline = item._p.xpath("./w:pPr/w:outlineLvl")
            level = int(match.group(1)) if match else (int(outline[0].get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")) + 1 if outline else None)
            if level and level <= 9:
                headings = headings[:level - 1]
                while len(headings) < level - 1:
                    headings.append("未命名上级章节")
                headings.append(text)
            result.append(block(text, "heading" if level else "paragraph", {"paragraph": pnum, "locator": f"段落 {pnum}", "heading_level": level}, headings))
        elif isinstance(item, Table):
            tnum += 1
            rows = []
            for number, row in enumerate(item.rows, 1):
                values = [""] * row.grid_cols_before + [cell_text(c.text) for c in row.cells] + [""] * row.grid_cols_after
                rows.append((number, values))
            result.extend(table_blocks(rows, {"table": tnum, "locator": f"表格 {tnum}"}, headings, config, f"docx-{tnum}"))
    return result, [], []


def extract_xlsx(path, config):
    values_book = load_workbook(path, data_only=True, read_only=False)
    formula_book = load_workbook(path, data_only=False, read_only=False)
    result, pages, warnings = [], [], []
    try:
        for sheet in values_book:
            if sheet.max_row * sheet.max_column > 500000:
                raise ValueError(f"工作表 {sheet.title} 超过 500000 个网格单元，请拆分后解析")
            merged = {}
            for region in sheet.merged_cells.ranges:
                value = sheet.cell(region.min_row, region.min_col).value
                for row in range(region.min_row, region.max_row + 1):
                    for col in range(region.min_col, region.max_col + 1):
                        merged[(row, col)] = value
            rows, missing_formula = [], 0
            for rnum, row in enumerate(sheet.iter_rows(), 1):
                values = []
                for cnum, cell in enumerate(row, 1):
                    value = merged.get((rnum, cnum), cell.value)
                    formula = formula_book[sheet.title].cell(rnum, cnum)
                    if value is None and formula.data_type == "f":
                        value = f"[公式未计算] {formula.value}"
                        missing_formula += 1
                    values.append(cell_text(value))
                if any(values):
                    rows.append((rnum, values))
            if missing_formula:
                warnings.append(f"工作表 {sheet.title} 有 {missing_formula} 个公式缺少缓存值，保留公式文本，不计算结果")
            result.extend(table_blocks(rows, {"sheet": sheet.title, "locator": f"工作表 {sheet.title}",
                "merged_ranges": [str(r) for r in sheet.merged_cells.ranges]}, [sheet.title], config, f"xlsx-{sheet.title}"))
            pages.append({"sheet": sheet.title, "status": "text" if rows else "empty", "rows": len(rows), "merged_ranges": len(sheet.merged_cells.ranges)})
    finally:
        values_book.close()
        formula_book.close()
    return result, pages, warnings


def extract_csv(path, config):
    text, encoding = decode_text(Path(path).read_bytes())
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    rows, previous_line = [], 0
    for values in reader:
        start = previous_line + 1
        previous_line = reader.line_num
        if any(v.strip() for v in values):
            rows.append((start, [cell_text(v) for v in values]))
    name = Path(path).name
    result = table_blocks(rows, {"sheet": "CSV", "encoding": encoding, "locator": "CSV 数据行"}, [], config, "csv-1")
    return result, [{"sheet": "CSV", "encoding": encoding, "rows": len(rows), "delimiter": dialect.delimiter}], []


def extract_txt(path):
    text, encoding = decode_text(Path(path).read_bytes())
    result, lines, start = [], [], 1
    for number, line in enumerate(text.splitlines(), 1):
        if line.strip():
            if not lines:
                start = number
            lines.append(clean(line))
        elif lines:
            result.append(block("\n".join(lines), "paragraph", {"line_start": start, "line_end": number - 1, "encoding": encoding, "locator": f"第 {start}–{number - 1} 行"}))
            lines = []
    if lines:
        result.append(block("\n".join(lines), "paragraph", {"line_start": start, "line_end": len(text.splitlines()), "encoding": encoding, "locator": f"第 {start} 行起"}))
    return result, [{"lines": len(text.splitlines()), "encoding": encoding}], []


def extract_pptx(path, config):
    presentation, result, pages = Presentation(path), [], []
    for snum, slide in enumerate(presentation.slides, 1):
        title = clean(slide.shapes.title.text) if slide.shapes.title else f"幻灯片 {snum}"
        section = [title]
        for shape_num, shape in enumerate(slide.shapes, 1):
            source = {"slide": snum, "shape": shape_num, "locator": f"幻灯片 {snum} / 对象 {shape_num}",
                "bbox_emu": [shape.left, shape.top, shape.width, shape.height]}
            if shape.has_text_frame:
                for pnum, paragraph in enumerate(shape.text_frame.paragraphs, 1):
                    text = clean(paragraph.text)
                    if text:
                        is_title = slide.shapes.title and shape.shape_id == slide.shapes.title.shape_id
                        result.append(block(text, "heading" if is_title else "paragraph", {**source, "paragraph": pnum}, section))
            if shape.has_table:
                rows = [(n, [cell_text(c.text) for c in row.cells]) for n, row in enumerate(shape.table.rows, 1)]
                result.extend(table_blocks(rows, source, section, config, f"ppt-{snum}-{shape_num}"))
        pages.append({"slide": snum, "title": title, "status": "text"})
    return result, pages, []


def ocr_page(page, language):
    # CLI execution is bounded; coordinates are converted back into PDF points.
    dpi = 200
    if page.rect.width * page.rect.height * (dpi / 72) ** 2 > 40000000:
        raise ValueError('页面尺寸过大，OCR 渲染超过 4000 万像素；请缩小或拆分页面')
    pix = page.get_pixmap(dpi=dpi, alpha=False)
    with tempfile.TemporaryDirectory(prefix="ocr-") as folder:
        image = Path(folder) / "page.png"
        pix.save(str(image))
        run = subprocess.run(["tesseract", str(image), "stdout", "-l", language, "--psm", "3", "tsv"],
            capture_output=True, text=True, encoding="utf-8", timeout=120)
        if run.returncode:
            raise RuntimeError(f"Tesseract: {run.stderr.strip()[:300]}")
        grouped = {}
        for word in csv.DictReader(io.StringIO(run.stdout), delimiter="\t", quoting=csv.QUOTE_NONE):
            if not word.get("text", "").strip():
                continue
            ident = (word["block_num"], word["par_num"], word["line_num"])
            grouped.setdefault(ident, []).append(word)
        lines = []
        for words in grouped.values():
            text = " ".join(w["text"] for w in words)
            x0 = min(int(w["left"]) for w in words)
            y0 = min(int(w["top"]) for w in words)
            x1 = max(int(w["left"]) + int(w["width"]) for w in words)
            y1 = max(int(w["top"]) + int(w["height"]) for w in words)
            confidence = sum(float(w["conf"]) for w in words) / len(words)
            lines.append((text, [v * 72 / dpi for v in [x0, y0, x1, y1]], round(confidence, 1)))
        return lines


def extract_pdf(path, config, progress):
    result, pages, warnings = [], [], []
    with fitz.open(path) as pdf:
        if pdf.needs_pass:
            raise ValueError("PDF 已加密，需要先移除密码")
        for number, page in enumerate(pdf, 1):
            status = {"page": number, "ocr": False, "status": "text"}
            try:
                texts = page.get_text("blocks", sort=True)
                usable = [b for b in texts if len(b) > 6 and b[6] == 0 and clean(b[4])]
                text = "".join(b[4] for b in usable)
                weak = len(text.strip()) < 20 or text.count("\ufffd") > max(1, len(text) * 0.2)
                needs_ocr = config["method"] == "ocr" or (config["method"] == "auto" and weak)
                if needs_ocr and config["ocrEnabled"]:
                    progress("ocr", round((number - 1) / max(1, len(pdf)) * 45), {"current_page": number})
                    lines = ocr_page(page, config["ocrLanguage"])
                    status["ocr"] = True
                    status["status"] = "ocr"
                    for text, bbox, confidence in lines:
                        result.append(block(text, "paragraph", {"page": number, "bbox": bbox, "ocr": True, "ocr_confidence": confidence, "locator": f"第 {number} 页 OCR"}))
                    if not lines:
                        raise ValueError("OCR 未识别到文字，页面可能为空白或分辨率不足")
                    status["mean_confidence"] = round(sum(x[2] for x in lines) / len(lines), 1)
                    if status["mean_confidence"] < 45:
                        status["warning"] = "OCR 平均置信度偏低，需要人工核对"
                else:
                    for b in usable:
                        result.append(block(clean(b[4]), "paragraph", {"page": number, "bbox": list(b[:4]), "ocr": False, "locator": f"第 {number} 页"}))
                    if weak:
                        status["warning"] = "页面可提取文本不足；文字模式或关闭 OCR 时未进行识别"
                        if not usable:
                            status["status"] = "unparsed"
                            status["reason"] = status["warning"]
            except Exception as error:
                status["status"] = "failed"
                status["reason"] = str(error)[:400]
            if status.get("warning") or status.get("reason"):
                warnings.append(f"第 {number} 页：{status.get('reason') or status.get('warning')}")
            pages.append(status)
            progress("extracting", round(number / max(1, len(pdf)) * 45), {"pages": pages, "warnings": warnings, "page_count": len(pdf)})
    return result, pages, warnings


def extract(path, extension, config, progress=lambda *_: None):
    path = Path(path)
    if path.stat().st_size > 512 * 1024 ** 2:
        raise ValueError("文件超过 512 MiB，已保存原件，请拆分后解析")
    progress("extracting", 0, {})
    if extension == ".pdf":
        return extract_pdf(path, config, progress)
    if extension == ".docx":
        return extract_docx(path, config)
    if extension == ".xlsx":
        return extract_xlsx(path, config)
    if extension == ".csv":
        return extract_csv(path, config)
    if extension == ".pptx":
        return extract_pptx(path, config)
    if extension == ".txt":
        return extract_txt(path)
    if extension in CONVERT:
        with tempfile.TemporaryDirectory(prefix="convert-") as folder:
            run = subprocess.run(["libreoffice", f"-env:UserInstallation=file://{folder}/profile", "--headless", "--convert-to", "pdf", "--outdir", folder, str(path)],
                capture_output=True, timeout=300)
            converted = Path(folder) / (path.stem + ".pdf")
            if run.returncode or not converted.exists():
                raise RuntimeError("LibreOffice 转换失败：" + run.stderr.decode(errors="replace")[:300])
            blocks, pages, warnings = extract_pdf(converted, config, progress)
            warnings.append("旧版 Office/OpenDocument 经 PDF 转换：来源页码是转换预览页码，不保证保留原表格结构")
            return blocks, pages, warnings
    raise ValueError("不支持此格式")


def section_prefix(section, config):
    return ("章节：" + " / ".join(section) + "\n") if section and config["sectionContext"] else ""


def source_for(parts, start=None, end=None):
    sources = [p["source"] for p in parts]
    first = dict(sources[0])
    first["positions"] = sources
    if start is not None:
        first["group_char_start"], first["group_char_end"] = start, end
    if first.get("page"):
        first["pages"] = sorted(set(s["page"] for s in sources if s.get("page")))
    if first.get("line_start"):
        first["line_end"] = sources[-1].get("line_end", first["line_start"])
    if first.get("row"):
        first["row_end"] = sources[-1].get("row", first["row"])
    return first


def make_chunk(content, kind, source, section, strategy, config):
    if len(content) > config["maxLength"]:
        label = {'qa':'问答对','table row':'表格数据行','table':'表格','text':'文本块'}.get(kind,kind)
        raise ValueError(f"{source.get('locator', '')} 的完整{label}超过最大长度 {config['maxLength']} 字符，请调大最大长度；未截断原始内容")
    return {"content": content, "content_type": kind, "source": source, "section_path": section,
        "strategy": strategy, "char_length": len(content)}


def text_chunks(parts, config, strategy):
    section = parts[0]["section_path"]
    prefix = section_prefix(section, config)
    capacity = config["maxLength"] - len(prefix)
    target = config["targetLength"] - len(prefix)
    if target < 20 or capacity < 20 or target <= config['overlapLength']:
        raise ValueError("章节路径占用过多长度，请增大分块长度或关闭章节路径补充")
    text, spans = "", []
    for part in parts:
        if text:
            text += "\n\n"
        start = len(text)
        text += part["content"]
        spans.append((start, len(text), part))
    offset, output, previous_end = 0, [], 0
    while offset < len(text):
        limit, aim = min(len(text), offset + capacity), min(len(text), offset + target)
        if limit == len(text) and (strategy == 'section' or len(text)-offset <= target):
            end = limit
        else:
            candidates = [m.end() + offset for m in re.finditer(r"\n\n|[。！？!?；;](?:\s|$)?|\.(?:\s|$)", text[offset:limit])]
            before = [n for n in candidates if offset + max(20, int(target * .55)) <= n <= aim]
            after = [n for n in candidates if aim < n <= limit]
            end = before[-1] if before else (after[0] if after else aim)
        affected = []
        for a, b, p in spans:
            if a >= end or b <= offset:
                continue
            local_start, local_end = max(0, offset-a), min(b-a,end-a)
            source = dict(p['source'], block_char_start=local_start, block_char_end=local_end)
            if 'line_start' in source:
                first_line = source['line_start']
                source['line_start'] = first_line + p['content'][:local_start].count('\n')
                source['line_end'] = first_line + p['content'][:max(local_start,local_end-1)].count('\n')
                source['locator'] = f"第 {source['line_start']}–{source['line_end']} 行"
            affected.append(dict(p, source=source))
        if not affected:
            affected = [parts[-1]]
        body = text[offset:end].strip()
        if body:
            source = source_for(affected, offset, end)
            source["overlap_mode"] = "characters_within_section"
            source['overlap_characters'] = max(0, previous_end-offset)
            output.append(make_chunk(prefix + body, "text", source, section, strategy, config))
        if end == len(text):
            break
        previous_end = end
        offset = max(offset + 1, end - config["overlapLength"])
    return output


def split_table(parts, config):
    section = parts[0]["section_path"]
    prefix = section_prefix(section, config)
    if config["headerContext"]:
        prefix += "表头：" + " | ".join(parts[0]["source"].get("column_names", [])) + "\n"
    output, current, fresh = [], [], False

    def emit(rows):
        source = source_for(rows)
        source['overlap_mode'] = 'whole_rows_within_table'
        previous_row_end = output[-1]['source'].get('row_end',0) if output else 0
        source['overlap_rows'] = sum(1 for row in rows if row['source'].get('row',0) <= previous_row_end)
        output.append(make_chunk(prefix + '\n'.join(p['content'] for p in rows), 'table', source, section, 'table', config))

    def tail(rows):
        retained, size = [], 0
        for p in reversed(rows):
            if size + len(p['content']) + 1 > config['overlapLength']:
                break
            retained.insert(0, p)
            size += len(p['content']) + 1
        return retained
    for part in parts:
        candidate = prefix + "\n".join(p["content"] for p in current + [part])
        if current and len(candidate) > config["maxLength"]:
            if fresh:
                emit(current)
            current, fresh = tail(current), False
            if len(prefix + "\n".join(p["content"] for p in current + [part])) > config["maxLength"]:
                current = []
        current.append(part)
        fresh = True
        if len(prefix + "\n".join(p["content"] for p in current)) > config["maxLength"]:
            make_chunk(prefix + part["content"], "table row", part["source"], section, "table", config)
        # Target is advisory; only flush early when the next row can remain whole.
        if len(prefix + "\n".join(p["content"] for p in current)) >= config["targetLength"]:
            emit(current)
            current, fresh = tail(current), False
    if current and fresh:
        emit(current)
    return output


def qa_chunks(parts, config, warnings):
    text = "\n".join(p["content"] for p in parts)
    spans, cursor = [], 0
    for part in parts:
        spans.append((cursor, cursor + len(part['content']), part))
        cursor += len(part['content']) + 1
    questions = list(re.finditer(r"(?m)^\s*(?:Q\s*\d*|问(?:题)?\s*\d*)\s*[:：]", text, re.I))
    if not questions:
        warnings.append("未识别到明确的 Q:/A: 或 问:/答: 标记，已回退为段落分块")
        return text_chunks(parts, config, "paragraph")
    output, section = [], parts[0]["section_path"]
    if text[:questions[0].start()].strip():
        intro = dict(parts[0], content=text[:questions[0].start()].strip())
        output.extend(text_chunks([intro], config, "paragraph"))
    for i, match in enumerate(questions):
        pair = text[match.start():questions[i + 1].start() if i + 1 < len(questions) else len(text)].strip()
        if not re.search(r"(?m)^\s*(?:A|答(?:案)?)\s*[:：]", pair, re.I):
            warnings.append("存在问题标记但未识别到答案标记，保留完整条目供人工核对")
        end = questions[i + 1].start() if i + 1 < len(questions) else len(text)
        affected = [p for a, b, p in spans if a < end and b > match.start()]
        output.append(make_chunk(section_prefix(section, config) + pair, "qa", source_for(affected, match.start(), end), section, "qa", config))
    return output


def chunk_blocks(blocks, config, warnings=None):
    warnings = warnings if warnings is not None else []
    result, group, key = [], [], None

    def flush():
        if not group:
            return
        if group[0]["content_type"] == "table":
            result.extend(split_table(group, config))
        elif config["strategy"] == "qa":
            result.extend(qa_chunks(group, config, warnings))
        else:
            result.extend(text_chunks(group, config, "section" if config["strategy"] == "section" else "paragraph"))

    for part in blocks:
        # Source line/page changes do not create groups. Sections/slides/sheets and
        # table identity do, so overlap never crosses into an unrelated section.
        next_key = (tuple(part["section_path"]), part["source"].get("table_id") if part["content_type"] == "table" else "text")
        if next_key != key:
            flush()
            group, key = [], next_key
        group.append(part)
    flush()
    for index, chunk in enumerate(result):
        chunk["chunk_index"] = index
    return result
