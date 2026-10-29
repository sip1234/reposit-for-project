"""Small real files for checking the original-preview/source-location contract."""
from pathlib import Path
import sys
import fitz
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

folder=Path(sys.argv[1])
folder.mkdir(parents=True, exist_ok=True)
pdf=fitz.open()
for i in range(2):
    pdf.new_page().insert_text((70,70), f"Source preview page {i+1}. Aeration process control and original text.")
pdf.save(folder/'source.pdf')
pdf.close()
doc=Document()
doc.add_heading('Operation manual',1)
doc.add_paragraph('Aeration control original paragraph with a real section path.')
table=doc.add_table(rows=2,cols=2)
for row,values in zip(table.rows,[('DO mg/L','Flow m3/h'),('2.3','40.5')]):
    for cell,value in zip(row.cells,values): cell.text=value
doc.save(folder/'source.docx')
book=Workbook()
book.active.title='Bioreactor'
book.active.append(['DO mg/L','Flow m3/h'])
book.active.append([2.3,40.5])
book.create_sheet('Effluent').append(['COD mg/L','Value'])
book['Effluent'].append(['Outlet',21])
book.save(folder/'source.xlsx')
slides=Presentation()
for i in range(2):
    slide=slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text=f'Process slide {i+1}'
    slide.placeholders[1].text='Original aeration control content.'
slides.save(folder/'source.pptx')
(folder/'source.csv').write_text('Point,DO mg/L,Description\nA01,0,"First line\nSecond line"\nA02,2.3,Normal\n',encoding='utf-8')
print('Created PDF, DOCX, XLSX, PPTX and multiline CSV preview fixtures')
