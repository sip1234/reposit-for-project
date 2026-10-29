"""Deterministic format fixtures, including an actual raster-only OCR page."""
import copy
import hashlib
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "worker"))
import fitz
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from structured_parser import DEFAULTS, block, extract, chunk_blocks


class StructuredParsingTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name)
        self.config = {**DEFAULTS, "targetLength": 160, "maxLength": 260, "overlapLength": 30}

    def tearDown(self):
        self.tmp.cleanup()

    def test_docx_paragraph_table_and_nested_headings(self):
        path = self.folder / "headings.docx"
        doc = Document()
        doc.add_heading("工艺规程", 1)
        doc.add_paragraph("这是完整的运行正文。")
        doc.add_heading("溶解氧", 2)
        table = doc.add_table(rows=3, cols=2)
        for row, values in zip(table.rows, [("DO (mg/L)", "流量 (m³/h)"), ("2.3", "0"), ("", "42.8")]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        doc.save(path)
        before = hashlib.sha256(path.read_bytes()).digest()
        blocks, _, _ = extract(path, ".docx", self.config)
        self.assertTrue(any(b["content"] == "这是完整的运行正文。" for b in blocks))
        rows = [b for b in blocks if b["content_type"] == "table"]
        self.assertEqual(rows[0]["section_path"], ["工艺规程", "溶解氧"])
        self.assertIn("DO (mg/L): 2.3", rows[0]["content"])
        self.assertIn("流量 (m³/h): 0", rows[0]["content"])
        self.assertIn("∅", rows[1]["content"])
        self.assertEqual(before, hashlib.sha256(path.read_bytes()).digest())

    def test_xlsx_sheets_headers_merges_zero_formula(self):
        path = self.folder / "sheets.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "生化池"
        sheet.append(["水质", None, "流量"])
        sheet.merge_cells("A1:B1")
        sheet.append(["DO mg/L", "NH4 mg/L", "m³/h"])
        sheet.append([0, None, 40.5])
        sheet.append(["=1+1", 3.2, 41])
        book.create_sheet("出水").append(["COD mg/L", "时间"])
        book["出水"].append([21, "2026-10-28"])
        book.save(path)
        config = {**self.config, "headerRows": 2}
        blocks, pages, warnings = extract(path, ".xlsx", config)
        self.assertEqual([p["sheet"] for p in pages], ["生化池", "出水"])
        self.assertIn("水质 / DO mg/L: 0", blocks[0]["content"])
        self.assertEqual(blocks[0]["source"]["row"], 3)
        self.assertIn("A1:B1", blocks[0]["source"]["merged_ranges"])
        self.assertTrue(any("公式" in w for w in warnings))

    def test_csv_encodings_quoted_rows_and_empty_cells(self):
        for encoding in ["utf-8-sig", "gb18030", "utf-16"]:
            with self.subTest(encoding=encoding):
                path = self.folder / "data.csv"
                path.write_bytes('测点,DO(mg/L),说明\nA01,0,"第一行\n第二行"\nA02,,正常\n'.encode(encoding))
                blocks, _, _ = extract(path, ".csv", self.config)
                self.assertIn("DO(mg/L): 0", blocks[0]["content"])
                self.assertIn("DO(mg/L): ∅", blocks[1]["content"])
                self.assertEqual(blocks[1]["source"]["row"], 4)

    def test_txt_short_lines_do_not_force_chunks(self):
        path = self.folder / "lines.txt"
        path.write_text("\n".join(["运行正常。"] * 60), encoding="utf-8")
        blocks, _, _ = extract(path, ".txt", self.config)
        self.assertEqual(len(blocks), 1)
        chunks = chunk_blocks(blocks, self.config)
        self.assertLess(len(chunks), 10)
        self.assertEqual(chunks[0]["source"]["line_start"], 1)
        self.assertEqual(chunks[-1]["source"]["line_end"], 60)

    def test_sections_overlap_unicode_and_length_bound(self):
        config = {**self.config, "sectionContext": False, "strategy": "section"}
        blocks = [block("甲😀。" * 150, "paragraph", {"paragraph": 1}, ["一"]), block("乙。" * 150, "paragraph", {"paragraph": 2}, ["二"])]
        chunks = chunk_blocks(blocks, config)
        self.assertTrue(all(len(c["content"]) == c["char_length"] <= config["maxLength"] for c in chunks))
        for left, right in zip(chunks, chunks[1:]):
            if left["section_path"] == right["section_path"]:
                self.assertEqual(left["content"][-30:], right["content"][:30])
            else:
                self.assertNotIn("甲", right["content"])
        self.assertEqual({tuple(c["section_path"]) for c in chunks}, {("一",), ("二",)})

    def test_section_strategy_keeps_short_chapter_whole(self):
        config = {**self.config,'sectionContext':False,'overlapLength':0}
        parts = [block('甲。'*45,'paragraph',{'paragraph':i},['同一章']) for i in range(2)]
        self.assertEqual(len(chunk_blocks(parts,{**config,'strategy':'section'})),1)
        self.assertEqual(len(chunk_blocks(parts,{**config,'strategy':'paragraph'})),2)

    def test_table_headers_repeat_and_overlap_whole_rows(self):
        config = {**self.config, "targetLength": 110, "maxLength": 160, "overlapLength": 45}
        parts = [block(f"DO(mg/L): {i} | 状态: 正常", "table", {"table_id": "one", "row": i + 2, "column_names": ["DO(mg/L)", "状态"]}) for i in range(20)]
        chunks = chunk_blocks(parts, config)
        self.assertGreater(len(chunks), 2)
        self.assertTrue(all(c["content"].startswith("表头：DO(mg/L)") for c in chunks))
        self.assertTrue(all(c["source"]["overlap_mode"] == "whole_rows_within_table" for c in chunks))
        self.assertLessEqual(chunks[1]["source"]["row"], chunks[0]["source"]["row_end"])
        self.assertEqual(chunks[-1]["source"]["row_end"], 21)

    def test_qa_pairs_complete_and_oversize_fails_explicitly(self):
        config = {**self.config, "strategy": "qa"}
        parts = [block("问：如何控制？\n答：调整曝气。\n\nQ: What is DO?\nA: Dissolved oxygen.", "paragraph", {"line_start": 1, "line_end": 5})]
        chunks = chunk_blocks(parts, config)
        self.assertEqual(len(chunks), 2)
        self.assertIn("答：调整曝气", chunks[0]["content"])
        self.assertIn("A: Dissolved", chunks[1]["content"])
        with self.assertRaisesRegex(ValueError, "超过最大长度"):
            chunk_blocks([block("Q: why?\nA: " + "x" * 300, "paragraph", {})], config)

    def test_pdf_text_and_real_scanned_page(self):
        path = self.folder / "mixed.pdf"
        source = fitz.open()
        page = source.new_page(width=600, height=200)
        page.insert_text((30, 65), "FLOW RATE 42.8 cubic metres per hour", fontsize=22)
        image = page.get_pixmap(dpi=200).tobytes("png")
        scanned = source.new_page(width=600, height=200)
        scanned.insert_image(scanned.rect, stream=image)
        source.new_page(width=600, height=200)
        source.save(path)
        source.close()
        blocks, pages, warnings = extract(path, ".pdf", self.config)
        self.assertEqual(len(pages), 3)
        self.assertFalse(pages[0]["ocr"])
        self.assertTrue(pages[1]["ocr"])
        self.assertTrue(any("FLOW" in b["content"] and b["source"]["page"] == 2 and b["source"]["bbox"] for b in blocks))
        self.assertEqual(pages[2]["status"], "failed")
        self.assertTrue(any("第 3 页" in warning for warning in warnings))
        _, disabled, _ = extract(path, ".pdf", {**self.config, "ocrEnabled": False})
        self.assertEqual(disabled[1]["status"], "unparsed")

    def test_pptx_slide_title_text_and_table_location(self):
        path = self.folder / "slides.pptx"
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = "运行情况"
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
        for row, values in zip(table.rows, [("DO mg/L", "流量"), ("2.3", "42.8")]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        deck.save(path)
        blocks, _, _ = extract(path, ".pptx", self.config)
        self.assertEqual(next(b for b in blocks if b['content']=='运行情况')['content_type'],'heading')
        row = next(b for b in blocks if b["content_type"] == "table")
        self.assertEqual(row["section_path"], ["运行情况"])
        self.assertEqual(row["source"]["slide"], 1)
        self.assertIn("DO mg/L: 2.3", row["content"])


if __name__ == "__main__":
    unittest.main()
